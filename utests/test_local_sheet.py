## { U-TEST

##
## === DEPENDENCIES
##

## stdlib
import datetime
import sys
import tempfile
import unittest
from pathlib import Path

## third-party
import openpyxl
from openpyxl.styles import Font

## local
from arxivscraper.support import local_sheet

##
## === TEST SUITE
##


class TestLocalSpreadsheet_Roundtrip(unittest.TestCase):

    def test_write_then_read_returns_same_rows(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            rows = [
                ["arxiv_id", "title"],
                ["0000.00000", "a title"],
                ["0001.00001", "another title"],
            ]
            spreadsheet.write_all_rows(rows)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=rows,
            )

    def test_read_missing_file_returns_empty(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            spreadsheet = local_sheet.LocalSpreadsheet(Path(temp_dir) / "papers.xlsx")
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[],
            )

    def test_multiline_and_commas_are_preserved(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path, sheet_name="Papers")
            abstract = "First paragraph.\nSecond paragraph with, a comma, and \"quotes\"."
            rows = [
                ["arxiv_id", "authors", "abstract"],
                ["0000.00000", "Author A, Author B", abstract],
            ]
            spreadsheet.write_all_rows(rows)
            restored = spreadsheet.read_all_rows()
            self.assertEqual(
                first=restored[1][2],
                second=abstract,
            )

    def test_write_creates_file_and_parent_directories(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "nested" / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([["arxiv_id"]])
            self.assertTrue(file_path.is_file())

    def test_write_creates_named_sheet_on_existing_workbook(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            other_sheet = local_sheet.LocalSpreadsheet(file_path, sheet_name="Other")
            other_sheet.write_all_rows([["col_a"], ["value"]])
            papers_sheet = local_sheet.LocalSpreadsheet(file_path, sheet_name="Papers")
            papers_sheet.write_all_rows([["arxiv_id"], ["0000.00000"]])
            workbook = openpyxl.load_workbook(file_path)
            self.assertIn(
                member="Other",
                container=workbook.sheetnames,
            )
            self.assertEqual(
                first=papers_sheet.read_all_rows(),
                second=[["arxiv_id"], ["0000.00000"]],
            )
            self.assertEqual(
                first=other_sheet.read_all_rows(),
                second=[["col_a"], ["value"]],
            )

    def test_rewrite_clears_previous_content(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0000.00000", "first"],
                ["0001.00001", "second"],
                ["0002.00002", "third"],
            ])
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0003.00003", "only"],
            ])
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[
                    ["arxiv_id", "title"],
                    ["0003.00003", "only"],
                ],
            )

    def test_read_empty_sheet_returns_no_rows(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            local_sheet.LocalSpreadsheet(file_path).write_all_rows([])
            self.assertEqual(
                first=local_sheet.LocalSpreadsheet(file_path).read_all_rows(),
                second=[],
            )

    def test_missing_sheet_name_returns_empty(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            local_sheet.LocalSpreadsheet(file_path, sheet_name="Papers").write_all_rows(
                [["arxiv_id"]],
            )
            self.assertEqual(
                first=local_sheet.LocalSpreadsheet(
                    file_path,
                    sheet_name="Other",
                ).read_all_rows(),
                second=[],
            )


class TestLocalSpreadsheet_Formatting(unittest.TestCase):
    """Formatting-only cells (styled but empty) must not widen/lengthen the grid."""

    def test_formatting_beyond_data_does_not_widen_rows(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            rows = [
                ["arxiv_id", "title", "authors"],
                ["0000.00000", "a", "b"],
            ]
            spreadsheet.write_all_rows(rows)
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            for row_index in range(2, 6):
                worksheet.cell(row=row_index, column=10).font = Font(bold=True)
            workbook.save(file_path)
            restored = spreadsheet.read_all_rows()
            self.assertTrue(all(len(row) == 3 for row in restored))
            self.assertEqual(
                first=restored,
                second=rows,
            )

    def test_formatting_empty_rows_below_does_not_add_rows(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0000.00000", "a"],
            ])
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            for column_index in range(1, 4):
                worksheet.cell(row=20, column=column_index).font = Font(bold=True)
            workbook.save(file_path)
            restored = spreadsheet.read_all_rows()
            self.assertEqual(
                first=restored,
                second=[
                    ["arxiv_id", "title"],
                    ["0000.00000", "a"],
                ],
            )

    def test_formatted_only_sheet_reads_as_empty(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([])
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            worksheet["C5"] = "   "
            worksheet["C5"].font = Font(bold=True)
            workbook.save(file_path)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[],
            )


class TestLocalSpreadsheet_FormattingPersistence(unittest.TestCase):
    """Rewrites must update content in place, keeping cell formatting intact."""

    def test_cell_styles_survive_rewrite(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0000.00000", "old-a"],
                ["0001.00001", "old-b"],
            ])
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            worksheet["A1"].font = Font(bold=True, color="FFFFFF")
            worksheet["B2"].font = Font(italic=True)
            workbook.save(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0002.00002", "new"],
            ])
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            self.assertTrue(worksheet["A1"].font.bold)
            self.assertEqual(
                first=worksheet["A1"].font.color.rgb,
                second="00FFFFFF",
            )
            self.assertTrue(worksheet["B2"].font.italic)
            self.assertEqual(
                first=worksheet["B2"].value,
                second="new",
            )

    def test_rewrite_clears_out_of_bounds_content_only(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0000.00000", "a"],
                ["0001.00001", "b"],
                ["0002.00002", "c"],
            ])
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0003.00003", "only"],
            ])
            workbook = openpyxl.load_workbook(file_path)
            worksheet = workbook["Papers"]
            self.assertEqual(
                first=worksheet["A1"].value,
                second="arxiv_id",
            )
            self.assertIsNone(worksheet["A3"].value)
            self.assertIsNone(worksheet["A4"].value)
            self.assertIsNone(worksheet["B4"].value)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[
                    ["arxiv_id", "title"],
                    ["0003.00003", "only"],
                ],
            )

    def test_rewrite_with_shorter_rows_blanks_leftover_cells(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            spreadsheet.write_all_rows([
                ["arxiv_id", "title", "authors"],
                ["0000.00000", "a", "someone"],
            ])
            spreadsheet.write_all_rows([
                ["arxiv_id", "title"],
                ["0000.00000", "a"],
            ])
            workbook = openpyxl.load_workbook(file_path)
            self.assertIsNone(workbook["Papers"]["C2"].value)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[
                    ["arxiv_id", "title"],
                    ["0000.00000", "a"],
                ],
            )


class TestLocalSpreadsheet_CellTypes(unittest.TestCase):
    def test_typed_cells_render_as_text(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.title = "Papers"
            worksheet.append([
                "0000.00000",
                7.5,
                7.0,
                datetime.datetime(2000, 1, 1, 0, 0),
            ])
            workbook.save(file_path)
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[["0000.00000", "7.5", "7", "2000-01-01"]],
            )

    def test_blank_cells_become_empty_strings(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "papers.xlsx"
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.title = "Papers"
            worksheet.append(["a", None, "", "d"])
            workbook.save(file_path)
            spreadsheet = local_sheet.LocalSpreadsheet(file_path)
            self.assertEqual(
                first=spreadsheet.read_all_rows(),
                second=[["a", "", "", "d"]],
            )


##
## === ENTRY POINT
##

if __name__ == "__main__":
    unittest.main()
    sys.exit(0)

## } U-TEST
